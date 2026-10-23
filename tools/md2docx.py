# -*- coding: utf-8 -*-
"""把论文 Markdown 转成 Word（.docx），图片按实际尺寸嵌入、表格转成 Word 表格。

用法：
    python tools/md2docx.py <输入.md> <输出.docx>

处理范围：标题（#/##/###/####）、正文、粗体/斜体/行内代码/上标引用、项目符号、
编号列表、管道表格、图片、引用块、分隔线。
"""

import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from PIL import Image


# ── 基础样式 ──────────────────────────────────────────────
def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.35

    for lvl, size in ((1, 16), (2, 14), (3, 12), (4, 11)):
        h = doc.styles["Heading {}".format(lvl)]
        h.font.name = "Times New Roman"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        h.paragraph_format.space_before = Pt(12 if lvl <= 2 else 9)
        h.paragraph_format.space_after = Pt(4)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for attr in ("left_margin", "right_margin"):
        setattr(sec, attr, Cm(2.5))
    sec.top_margin = sec.bottom_margin = Cm(2.4)


# ── 行内格式：**粗** *斜* `码` <sup>x</sup> ─────────────────
TOKEN = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`|<sup>.+?</sup>)")


def add_inline(par, text):
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif piece.startswith("<sup>") and piece.endswith("</sup>"):
            r = par.add_run(piece[5:-6])
            r.font.superscript = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


# ── 表格 ─────────────────────────────────────────────────
def split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, h)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
    for i, row in enumerate(body, start=1):
        for j in range(ncol):
            cell = t.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, row[j] if j < len(row) else "")
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


# ── 图片 ─────────────────────────────────────────────────
MAX_W_CM = 16.0


def add_image(doc, path):
    if not os.path.exists(path):
        doc.add_paragraph("[缺图: {}]".format(path))
        return
    im = Image.open(path)
    dpi = im.info.get("dpi", (300, 300))
    w_cm = im.size[0] / float(dpi[0]) * 2.54
    h_cm = im.size[1] / float(dpi[1]) * 2.54
    if w_cm > MAX_W_CM:                       # 超出版心则等比缩小
        ratio = MAX_W_CM / w_cm
        w_cm, h_cm = MAX_W_CM, h_cm * ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(w_cm))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)


# ── 主流程 ───────────────────────────────────────────────
def convert(src, dst):
    base = os.path.dirname(os.path.abspath(src))
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    setup(doc)

    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # 图片（图注在下一段，由后续文本处理）
        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s)
        if m:
            add_image(doc, os.path.normpath(os.path.join(base, m.group(1))))
            i += 1
            continue

        # 表格：连续以 | 开头的行
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = [split_row(s)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # 标题
        m = re.match(r"^(#{1,4}) (.+)$", s)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue

        # 分隔线
        if re.match(r"^-{3,}$", s) or s == "***":
            i += 1
            continue

        # 引用块
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            for b in buf:
                if not b:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, b)
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # 无序列表
        if re.match(r"^[-*] ", s):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, s[2:])
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\d+)\. (.+)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # 空行
        if not s:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, s)
        p.paragraph_format.first_line_indent = Cm(0)
        i += 1

    doc.save(dst)
    return doc


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    convert(sys.argv[1], sys.argv[2])
    print("已生成:", sys.argv[2], os.path.getsize(sys.argv[2]), "字节")
